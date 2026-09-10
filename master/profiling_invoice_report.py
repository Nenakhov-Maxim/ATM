import os
from copy import copy
from collections import OrderedDict, defaultdict
from datetime import datetime, time

from django.db.models import Prefetch, Q
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from .forms import REPORT_TIME_ZONE
from .models import HistoryProfileRecords, OffsShtrips, Tasks


SHIFT_1 = 'shift_1'
SHIFT_2 = 'shift_2'
SHIFT_3 = 'shift_3'
SHIFT_KEYS = (SHIFT_1, SHIFT_2, SHIFT_3)


def create_profiling_invoice_report(start_date, end_date, user, output_filename=None):
    if not os.path.exists('excel_files'):
        os.makedirs('excel_files')

    if output_filename is None:
        output_filename = f'Накладная на линию профилирования от {datetime.now().date()}.xlsx'

    filepath = os.path.join('excel_files', output_filename)
    tasks = _get_completed_tasks(start_date, end_date, user)
    rows = _build_report_rows(tasks)
    _write_workbook(filepath, rows, start_date, end_date)
    return filepath


def _get_completed_tasks(start_date, end_date, user):
    profile_records = HistoryProfileRecords.objects.filter(
        created_at__gte=start_date,
        created_at__lt=end_date,
    ).select_related('user').order_by('created_at', 'id')
    shtrips_records = OffsShtrips.objects.filter(
        created_at__gte=start_date,
        created_at__lt=end_date,
    ).select_related('type_value_id').order_by('created_at', 'id')

    return Tasks.objects.filter(
        Q(task_status_id=2) &
        Q(task_timedate_end_fact__gte=start_date) &
        Q(task_timedate_end_fact__lt=end_date) &
        Q(
            Q(production_area=user.production_area_id) |
            Q(task_workplace__production_area_id=user.production_area_id)
        )
    ).select_related(
        'task_profile_type',
        'task_coating_type',
    ).prefetch_related(
        Prefetch(
            'history_profile_records',
            queryset=profile_records,
            to_attr='invoice_profile_records',
        ),
        Prefetch(
            'history_offs_shtrips',
            queryset=shtrips_records,
            to_attr='invoice_shtrips_records',
        ),
    ).order_by('task_profile_type__profile_name', 'task_profile_length', 'id')


def _build_report_rows(tasks):
    profile_groups = OrderedDict()
    for task in tasks:
        if not any(record.amount for record in task.invoice_profile_records):
            continue
        coating_rows = {}

        def rows_for(coating):
            if coating is None:
                coating = task.task_coating_thickness or ''
            if coating not in coating_rows:
                snapshot = copy(task)
                snapshot.task_coating_thickness = coating
                profile_key = _profile_group_key(snapshot)
                group = profile_groups.setdefault(profile_key, _empty_profile_group(snapshot))
                detail_key = _task_group_key(snapshot)
                detail = group['details'].setdefault(detail_key, _empty_detail_row(snapshot))
                coating_rows[coating] = (group, detail)
            return coating_rows[coating]

        for record in task.invoice_profile_records:
            if not record.amount:
                continue
            group, detail = rows_for(record.coating_thickness)
            shift = _shift_key(record.created_at)
            group['profile_by_shift'][shift] += record.amount
            detail['profile_by_shift'][shift] += record.amount

        for record in task.invoice_shtrips_records:
            group, detail = rows_for(record.coating_thickness)
            shift = _shift_key(record.created_at)
            group['shtrips_by_shift'][shift] += _shtrips_weight_kg(record)
            detail['shtrips_by_shift'][shift] += _shtrips_weight_kg(record)

    return list(profile_groups.values())


def _empty_profile_group(task):
    return {
        'order_number': task.order_label,
        'profile_name': _profile_report_name(task),
        'shtrips_name': _shtrips_report_name(task),
        'profile_by_shift': defaultdict(float),
        'shtrips_by_shift': defaultdict(float),
        'details': OrderedDict(),
    }


def _empty_detail_row(task):
    return {
        'order_number': task.order_label,
        'coating': _detail_coating_label(task),
        'profile_length': task.task_profile_length or 0,
        'profile_by_shift': defaultdict(float),
        'shtrips_by_shift': defaultdict(float),
    }


def _profile_group_key(task):
    return (
        task.task_order_number,
        task.task_profile_type_id,
        task.task_profile_material,
        task.task_coating_thickness,
        task.task_coating_type_id,
    )


def _task_group_key(task):
    return (
        task.task_order_number,
        task.task_profile_type_id,
        round(float(task.task_profile_length or 0), 3),
        task.task_coating_type_id,
        task.task_coating_thickness,
        task.task_coating_area,
    )


def _profile_report_name(task):
    parts = [task.task_profile_type.profile_name]
    material_thickness = _format_number(task.task_profile_material)
    coating_thickness = task.task_coating_thickness or ''
    coating_type = str(task.task_coating_type) if task.task_coating_type else ''
    suffix = f'{material_thickness}{coating_thickness}{coating_type}'
    if suffix:
        parts.append(suffix)
    return ' '.join(parts)


def _detail_coating_label(task):
    return str(task.task_coating_type) if task.task_coating_type else ''


def _shtrips_report_name(task):
    shtrips_name = task.task_profile_type.association_name_shtrips or ''
    material_thickness = _format_number(task.task_profile_material)
    coating_thickness = task.task_coating_thickness or ''

    first_line = shtrips_name
    if material_thickness:
        first_line = f'{first_line}х{material_thickness}' if first_line else f'х{material_thickness}'
    if coating_thickness:
        return f'{first_line}\n{coating_thickness}'
    return first_line


def _format_number(value):
    if value is None:
        return ''
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).replace('.', ',')


def _profile_amounts_by_shift(task):
    amounts = defaultdict(float)
    for record in task.invoice_profile_records:
        amounts[_shift_key(record.created_at)] += record.amount
    return amounts


def _shtrips_by_shift(task):
    amounts = defaultdict(float)
    for shtrips in task.invoice_shtrips_records:
        amounts[_shift_key(shtrips.created_at)] += _shtrips_weight_kg(shtrips)
    return amounts


def _shtrips_weight_kg(shtrips):
    return shtrips.value


def _shift_key(value):
    local_value = timezone.localtime(value, REPORT_TIME_ZONE) if timezone.is_aware(value) else value
    current_time = local_value.time()

    if time(8, 0) <= current_time < time(17, 0):
        return SHIFT_1
    if current_time >= time(17, 0):
        return SHIFT_2
    if current_time < time(1, 0):
        return SHIFT_2
    return SHIFT_3


def _write_workbook(filepath, rows, start_date, end_date):
    wb = Workbook()
    ws = wb.active
    ws.title = 'Накладная'

    header_font = Font(bold=True, size=10)
    profile_font = Font(bold=True, size=10)
    detail_font = Font(size=10)
    thin = Side(border_style='thin', color='000000')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal='center', vertical='center', wrap_text=True)
    left = Alignment(horizontal='left', vertical='center', wrap_text=True)
    detail_left = Alignment(horizontal='left', vertical='center')

    _setup_example_columns(ws)
    _write_example_header(ws, header_font, border, center)

    current_row = 4
    for profile_group in rows:
        _write_profile_group_row(ws, current_row, profile_group, profile_font, border, center, left)
        current_row += 1

        for detail_row in profile_group['details'].values():
            _write_detail_row(ws, current_row, detail_row, detail_font, border, center, detail_left)
            current_row += 1

    _write_total_row(ws, current_row, rows, border, center)
    wb.save(filepath)


def _setup_example_columns(ws):
    for column in range(1, 26):
        ws.column_dimensions[ws.cell(1, column).column_letter].width = 6

    ws.column_dimensions['A'].width = 14
    ws.column_dimensions['B'].width = 12
    ws.column_dimensions['C'].width = 5
    ws.column_dimensions['D'].width = 4
    ws.column_dimensions['E'].width = 8


def _write_example_header(ws, header_font, border, center):
    merges = [
        'A1:A3', 'B1:E3', 'F1:G1', 'H1:I1', 'J1:K1', 'L1:M2', 'N1:O2',
        'P1:Q3', 'R1:S1', 'T1:U1', 'V1:W1', 'X1:Y2',
        'F2:G2', 'H2:I2', 'J2:K2', 'R2:S2', 'T2:U2', 'V2:W2',
        'F3:G3', 'H3:I3', 'J3:K3', 'L3:M3', 'N3:O3',
        'R3:S3', 'T3:U3', 'V3:W3', 'X3:Y3',
    ]
    for cells in merges:
        ws.merge_cells(cells)

    values = {
        'A1': '№ заказа',
        'B1': 'Профиль',
        'F1': '1 см.',
        'H1': '2см.',
        'J1': '3см.',
        'L1': 'Всего',
        'N1': 'Всего',
        'P1': 'Штрипс',
        'R1': '1см.',
        'T1': '2см.',
        'V1': '3см.',
        'X1': 'Всего',
        'F2': 'Кол-во',
        'H2': 'Кол-во ',
        'J2': 'Кол-во',
        'R2': 'К-во штр.',
        'T2': 'К-во штр.',
        'V2': 'К-во штр.',
        'F3': '(штук)',
        'H3': '(штук)',
        'J3': '(штук)',
        'L3': '(штук)',
        'N3': '(п/м)',
        'R3': '(кг)',
        'T3': '(кг)',
        'V3': '(кг)',
        'X3': 'м/п',
    }
    for cell, value in values.items():
        ws[cell] = value

    for row in ws.iter_rows(min_row=1, max_row=3, min_col=1, max_col=25):
        for cell in row:
            cell.font = header_font
            cell.border = border
            cell.alignment = center


def _write_profile_group_row(ws, row_number, profile_group, font, border, center, left):
    _merge_data_row(ws, row_number)
    profile_by_shift = profile_group['profile_by_shift']
    shtrips_by_shift = profile_group['shtrips_by_shift']

    ws.cell(row_number, 1).value = profile_group['order_number']
    ws.cell(row_number, 2).value = profile_group['profile_name']
    ws.cell(row_number, 6).value = _number_or_empty(profile_by_shift[SHIFT_1])
    ws.cell(row_number, 8).value = _number_or_empty(profile_by_shift[SHIFT_2])
    ws.cell(row_number, 10).value = _number_or_empty(profile_by_shift[SHIFT_3])
    ws.cell(row_number, 12).value = _number_or_empty(sum(profile_by_shift.values()))
    ws.cell(row_number, 16).value = profile_group['shtrips_name']
    ws.cell(row_number, 18).value = _number_or_empty(round(shtrips_by_shift[SHIFT_1], 2))
    ws.cell(row_number, 20).value = _number_or_empty(round(shtrips_by_shift[SHIFT_2], 2))
    ws.cell(row_number, 22).value = _number_or_empty(round(shtrips_by_shift[SHIFT_3], 2))
    ws.cell(row_number, 24).value = _number_or_empty(round(sum(shtrips_by_shift.values()), 2))

    _style_data_row(ws, row_number, font, border, center, left)


def _write_detail_row(ws, row_number, detail_row, font, border, center, left):
    _merge_data_row(ws, row_number)
    profile_by_shift = detail_row['profile_by_shift']
    total_amount = sum(profile_by_shift.values())
    total_length = round(total_amount * float(detail_row['profile_length'] or 0), 2)
    shtrips_by_shift = detail_row['shtrips_by_shift']

    ws.cell(row_number, 1).value = detail_row['order_number']
    ws.cell(row_number, 2).value = detail_row['coating']
    ws.cell(row_number, 4).value = 'L='
    ws.cell(row_number, 5).value = _format_number(detail_row['profile_length'])
    ws.cell(row_number, 6).value = _number_or_empty(profile_by_shift[SHIFT_1])
    ws.cell(row_number, 8).value = _number_or_empty(profile_by_shift[SHIFT_2])
    ws.cell(row_number, 10).value = _number_or_empty(profile_by_shift[SHIFT_3])
    ws.cell(row_number, 12).value = _number_or_empty(total_amount)
    ws.cell(row_number, 14).value = _number_or_empty(total_length)
    ws.cell(row_number, 18).value = _number_or_empty(round(shtrips_by_shift[SHIFT_1], 2))
    ws.cell(row_number, 20).value = _number_or_empty(round(shtrips_by_shift[SHIFT_2], 2))
    ws.cell(row_number, 22).value = _number_or_empty(round(shtrips_by_shift[SHIFT_3], 2))
    ws.cell(row_number, 24).value = _number_or_empty(round(sum(shtrips_by_shift.values()), 2))

    _style_data_row(ws, row_number, font, border, center, left)


def _write_total_row(ws, row_number, rows, border, center):
    _merge_data_row(ws, row_number)
    total_length = 0
    total_shtrips = 0
    for profile_group in rows:
        total_shtrips += sum(profile_group['shtrips_by_shift'].values())
        for detail_row in profile_group['details'].values():
            total_amount = sum(detail_row['profile_by_shift'].values())
            total_length += total_amount * float(detail_row['profile_length'] or 0)

    ws.cell(row_number, 14).value = _number_or_empty(round(total_length, 2))
    ws.cell(row_number, 24).value = _number_or_empty(round(total_shtrips, 2))
    _style_data_row(ws, row_number, Font(bold=True, size=10), border, center, center)


def _merge_data_row(ws, row_number):
    for start_column in range(6, 26, 2):
        ws.merge_cells(
            start_row=row_number,
            start_column=start_column,
            end_row=row_number,
            end_column=start_column + 1,
        )


def _style_data_row(ws, row_number, font, border, center, left):
    for cell in ws[row_number]:
        cell.font = font
        cell.border = border
        cell.alignment = left if cell.column in (1, 2, 4) else center


def _number_or_empty(value):
    if value is None:
        return ''
    if value == 0:
        return ''
    if isinstance(value, float) and value.is_integer():
        return int(value)
    return value
