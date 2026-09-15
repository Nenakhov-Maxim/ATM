from datetime import date, datetime, timedelta, timezone
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch

from django.test import RequestFactory, SimpleTestCase, TestCase
from openpyxl import load_workbook

from login.models import ProductionArea, User, Workplace
from .completed_work_report import completed_work_rows, create_completed_work_report
from .forms import ReportForm
from .models import ProfileType, ShtripsValueType, Tasks, TaskStatus
from .profiling_invoice_report import _build_report_rows, _get_completed_tasks, create_profiling_invoice_report
from .report_periods import report_filename
from .shifts import shift_bounds
from .views import new_report, profiling_invoice_report


class ReportSelectionFormTests(SimpleTestCase):
    def test_disjoint_shifts_are_sorted_unique_and_ignore_inactive_dates(self):
        form = ReportForm({'report_mode': 'shifts', 'production_date': '2026-12-31',
                           'report_shifts': ['3', '1', '3'], 'date_start': 'invalid'})
        self.assertTrue(form.is_valid(), form.errors)
        self.assertEqual(form.cleaned_data['report_shifts'], [1, 3])
        self.assertEqual(form.cleaned_data['periods'], [shift_bounds(date(2026, 12, 31), i) for i in (1, 3)])
        self.assertEqual(form.cleaned_data['date_end'], datetime(2027, 1, 1, 3, tzinfo=timezone.utc))
        self.assertEqual(form.cleaned_data['period_label'], '31.12.2026, смены 1, 3')
        self.assertNotEqual(report_filename('Акт', form.cleaned_data), report_filename('Акт', form.cleaned_data))

    def test_missing_and_invalid_shift_fields(self):
        for changes in ({'report_shifts': []}, {'report_shifts': ['4']}, {'production_date': ''}, {'production_date': 'bad'}):
            data = {'report_mode': 'shifts', 'production_date': '2026-09-15', 'report_shifts': ['1']}
            data.update(changes)
            with self.subTest(changes=changes):
                self.assertFalse(ReportForm(data).is_valid())

    def test_legacy_period_ignores_inactive_shift_fields(self):
        form = ReportForm({'date_start': '2026-09-15T08:00', 'date_end': '2026-09-16T08:00',
                           'production_date': 'bad', 'report_shifts': ['4']})
        self.assertTrue(form.is_valid(), form.errors)
        self.assertEqual(len(form.cleaned_data['periods']), 1)
        self.assertEqual(form.cleaned_data['periods'][0][1] - form.cleaned_data['periods'][0][0], timedelta(days=1))

    def test_empty_equal_and_unknown_periods_are_invalid(self):
        for data in ({}, {'date_start': '2026-09-15T08:00', 'date_end': '2026-09-15T08:00'}, {'report_mode': 'bad'}):
            with self.subTest(data=data):
                self.assertFalse(ReportForm(data).is_valid())


class ReportSelectionDataTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.day = date(2026, 9, 15)
        cls.periods = [shift_bounds(cls.day, i) for i in (1, 3)]
        cls.area = ProductionArea.objects.create(production_area_name='Area')
        cls.master = User.objects.create(username='master', is_superuser=True, production_area_id=cls.area)
        cls.workers = [User.objects.create(username=f'worker-{i}', first_name='Same', last_name='Name') for i in range(2)]
        cls.line = Workplace.objects.create(workplace_name='Line', type_of_equipment='Line', inv_number=1, production_area_id=cls.area)
        cls.profile = ProfileType.objects.create(profile_name='Profile', association_name_shtrips='240')
        for pk in (1, 2):
            TaskStatus.objects.create(pk=pk, status_name=str(pk))
        cls.kg = ShtripsValueType.objects.create(pk=1, type_offs_shtrips='kg')

    def task(self, **changes):
        fields = dict(task_name='Task', task_profile_amount=1000, task_profile_type=self.profile,
                      task_profile_length=3, task_workplace=self.line, task_status_id=2,
                      task_timedate_end_fact=self.periods[1][0] + timedelta(hours=1))
        fields.update(changes)
        return Tasks.objects.create(**fields)

    def seed_records(self, task):
        # Include every boundary: 08:00, 17:00, 01:00 and next day's 08:00.
        for index, timestamp in enumerate((self.periods[0][0], self.periods[0][1], self.periods[1][0], self.periods[1][1])):
            task.history_profile_records.create(user=self.workers[index % 2], amount=(index + 1) * 10, created_at=timestamp)
            task.history_offs_shtrips.create(value=(index + 1) * 100, type_value_id=self.kg, created_at=timestamp)

    def test_invoice_disjoint_periods_filter_output_and_strips(self):
        task = self.task(task_shift=2, task_shift_date=self.day)
        self.seed_records(task)
        with self.assertNumQueries(3):
            rows = _build_report_rows(_get_completed_tasks(self.periods[0][0], self.periods[-1][1], self.master, self.periods))
        self.assertEqual(dict(rows[0]['profile_by_shift']), {'shift_1': 10, 'shift_3': 30})
        self.assertEqual(dict(rows[0]['shtrips_by_shift']), {'shift_1': 100, 'shift_3': 300})

    def test_invoice_preserves_completed_task_window_and_area_rules(self):
        included = self.task()
        self.task(task_status_id=1)
        self.task(task_timedate_end_fact=self.periods[0][1])
        self.task(task_timedate_end_fact=self.periods[1][1])
        other_area = ProductionArea.objects.create(production_area_name='Other')
        self.task(task_workplace=None, production_area=other_area)
        selected = _get_completed_tasks(self.periods[0][0], self.periods[-1][1], self.master, self.periods)
        self.assertEqual(list(selected.values_list('id', flat=True)), [included.pk])

    def test_act_uses_actual_output_not_task_last_update_or_planned_shift(self):
        task = self.task(task_status_id=1, task_shift=2, task_shift_date=self.day)
        self.seed_records(task)
        Tasks.objects.filter(pk=task.pk).update(last_update=self.periods[-1][1] + timedelta(days=5))
        with self.assertNumQueries(3):
            rows = completed_work_rows(self.periods, self.master)
        self.assertEqual(rows[task.pk]['data'][0][3], 120)
        self.assertIn('100.0; 300.0', rows[task.pk]['label'])
        self.assertNotIn('200.0', rows[task.pk]['label'])

    def test_act_keeps_same_name_workers_separate_and_respects_area(self):
        task = self.task()
        for user, amount in zip(self.workers, (10, 20)):
            task.history_profile_records.create(user=user, amount=amount, created_at=self.periods[0][0])
        other_area = ProductionArea.objects.create(production_area_name='Other')
        other = self.task(task_workplace=None, production_area=other_area)
        self.seed_records(other)
        rows = completed_work_rows(self.periods, self.master)
        self.assertEqual(list(rows), [task.pk])
        self.assertEqual([row[3] for row in rows[task.pk]['data']], [30, 60])

    def test_empty_selection_returns_no_rows(self):
        self.seed_records(self.task())
        self.assertEqual(completed_work_rows([], self.master), {})
        self.assertFalse(_get_completed_tasks(self.periods[0][0], self.periods[-1][1], self.master, []).exists())

    def test_all_shifts_match_full_day_and_single_second_shift_is_isolated(self):
        task = self.task()
        self.seed_records(task)
        shifts = [shift_bounds(self.day, i) for i in (1, 2, 3)]
        full_day = [(shifts[0][0], shifts[-1][1])]
        self.assertEqual(completed_work_rows(shifts, self.master), completed_work_rows(full_day, self.master))
        selected = _build_report_rows(_get_completed_tasks(shifts[0][0], shifts[-1][1], self.master, shifts))
        legacy = _build_report_rows(_get_completed_tasks(shifts[0][0], shifts[-1][1], self.master))
        self.assertEqual(selected, legacy)
        second_shift = completed_work_rows([shifts[1]], self.master)
        self.assertEqual(second_shift[task.pk]['data'][0][3], 60)

    def test_both_workbooks_include_selection_and_filtered_totals(self):
        self.seed_records(self.task())
        label = '15.09.2026, смены 1, 3'
        with TemporaryDirectory() as directory:
            act = create_completed_work_report(self.periods, self.master, str(Path(directory) / 'act.xlsx'), label)
            book = load_workbook(act)
            self.assertEqual(book.active['B1'].value, label)
            self.assertEqual(book.active['D5'].value, 120)
            book.close()
            invoice = create_profiling_invoice_report(self.periods[0][0], self.periods[-1][1], self.master,
                                                      str(Path(directory) / 'invoice.xlsx'), self.periods, label)
            book = load_workbook(invoice)
            self.assertEqual(book.properties.description, label)
            self.assertEqual(book.active['F4'].value, 10)
            self.assertIn(book.active['H4'].value, (None, 0))
            self.assertEqual(book.active['J4'].value, 30)
            self.assertEqual(book.active['L4'].value, 40)
            book.close()

    def test_endpoints_reject_invalid_input_and_get(self):
        for view in (new_report, profiling_invoice_report):
            for method, data, status in (('get', {}, 405), ('post', {'report_mode': 'shifts', 'production_date': '2026-09-15'}, 400)):
                request = getattr(RequestFactory(), method)('/report/', data)
                request.user = self.master
                self.assertEqual(view(request).status_code, status)

    def test_endpoints_pass_disjoint_ranges_to_writers(self):
        for view, writer in ((new_report, 'create_completed_work_report'), (profiling_invoice_report, 'create_profiling_invoice_report')):
            request = RequestFactory().post('/report/', {'report_mode': 'shifts', 'production_date': '2026-09-15', 'report_shifts': ['1', '3']})
            request.user = self.master
            with TemporaryDirectory() as directory:
                output = Path(directory) / 'report.xlsx'
                output.touch()
                with patch(f'master.views.{writer}', return_value=str(output)) as generate:
                    response = view(request)
                    self.assertEqual(response.status_code, 200)
                    self.assertIn('attachment;', response['Content-Disposition'])
                    self.assertEqual(generate.call_args.args[0] if view == new_report else generate.call_args.kwargs['periods'], self.periods)
                    response.close()
